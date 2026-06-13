"""
失败客户分析数据分析模板
==========================

输出：tools/失败客户分析模板.xlsx

工作表：
    1. 使用说明
    2. 失败客户登记表（30+ 客户主表）
    3. 7 大失败类型分布
    4. 销售个人失败分析
    5. 产品 / 价格失败分析
    6. 流程改进追踪
    7. 月度复盘汇总
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule


THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PRIMARY = "1F4E79"
LIGHT_BG = "DEEBF7"
INPUT_BG = "FFF2CC"
SUCCESS = "E2EFDA"
WARN = "FCE4D6"
CRITICAL = "FFB6C1"
WHITE = "FFFFFF"

H_FILL = PatternFill("solid", fgColor=PRIMARY)
H_FONT = Font(name="Microsoft YaHei", size=11, bold=True, color=WHITE)
TITLE_FONT = Font(name="Microsoft YaHei", size=16, bold=True, color=PRIMARY)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def style_h(c):
    c.fill = H_FILL; c.font = H_FONT; c.alignment = CENTER; c.border = BORDER


def style_t(c, fill=None, align=None, money=False, pct=False, bold=False):
    c.font = Font(name="Microsoft YaHei", size=10, bold=bold)
    c.alignment = align or CENTER; c.border = BORDER
    if fill: c.fill = PatternFill("solid", fgColor=fill)
    if money: c.number_format = '"¥"#,##0'
    if pct: c.number_format = "0.0%"


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_readme(ws):
    ws.title = "使用说明"
    widths(ws, [4, 26, 70])
    ws.merge_cells("B2:C2")
    ws["B2"] = "失败客户分析数据模板"
    ws["B2"].font = TITLE_FONT
    ws["B2"].alignment = LEFT

    rows = [
        ("适用对象", "销售总监 + CS 总监 + 老板"),
        ("使用步骤",
         "① 每个失败客户填『失败客户登记表』黄色单元格\n"
         "② 切换『7 大类型分布』自动统计\n"
         "③ 切换『销售个人失败分析』看哪个销售问题大\n"
         "④ 切换『产品价格分析』看产品问题\n"
         "⑤ 切换『流程改进追踪』管理改进动作\n"
         "⑥ 切换『月度复盘』看趋势"),
        ("关键概念",
         "7 大失败类型：\n"
         "A. 时机型（25%）- 不是时候\n"
         "B. 价格型（20%）- 嫌贵\n"
         "C. 价值型（15%）- 没听懂\n"
         "D. 关系型（12%）- 决策人变\n"
         "E. 信任型（10%）- 怕失败\n"
         "F. 竞品型（8%）- 竞品锁\n"
         "G. 内部型（10%）- 团队反对"),
        ("分析频次",
         "每周 1 次：销售总监审核新增失败客户\n"
         "每月 1 次：老板 + 全员复盘\n"
         "每季度 1 次：产品 + 销售 + CS 三方深度复盘"),
    ]
    for i, (k, v) in enumerate(rows, start=4):
        c1 = ws.cell(row=i, column=2, value=k); style_t(c1, fill=PRIMARY, bold=True)
        c1.font = Font(name="Microsoft YaHei", size=11, bold=True, color=WHITE)
        c2 = ws.cell(row=i, column=3, value=v); style_t(c2, align=LEFT)
        ws.row_dimensions[i].height = max(60, len(v) * 0.8)


def build_registry(ws):
    ws.title = "失败客户登记表"
    widths(ws, [5, 8, 22, 14, 14, 8, 12, 12, 14, 30, 14, 30, 12])

    ws.merge_cells("A1:M1")
    ws["A1"] = "失败客户登记表（黄色=输入，绿色=自动）"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = LEFT

    headers = [
        "#", "段", "客户名", "负责销售", "失败日期",
        "失败类型", "接触次数", "拜访次数", "最大反对意见",
        "销售自评（错的事）", "客户访谈日期", "客户原话",
        "处置"
    ]
    for j, h in enumerate(headers, 1):
        style_h(ws.cell(row=3, column=j, value=h))
    ws.row_dimensions[3].height = 32

    # 30 行
    for i in range(4, 34):
        idx = i - 3
        style_t(ws.cell(row=i, column=1, value=idx), fill=LIGHT_BG, bold=True)
        for col in range(2, 14):
            style_t(ws.cell(row=i, column=col), fill=INPUT_BG, align=LEFT)

    # 数据有效性
    dv_seg = DataValidation(type="list", formula1='"S,A,B,C"', allow_blank=True)
    ws.add_data_validation(dv_seg); dv_seg.add("B4:B33")

    dv_type = DataValidation(type="list",
        formula1='"A时机,B价格,C价值,D关系,E信任,F竞品,G内部"', allow_blank=True)
    ws.add_data_validation(dv_type); dv_type.add("F4:F33")

    dv_action = DataValidation(type="list",
        formula1='"6月再启动,降低期望换签,彻底放弃"', allow_blank=True)
    ws.add_data_validation(dv_action); dv_action.add("M4:M33")

    # 销售下拉
    dv_sales = DataValidation(type="list",
        formula1='"A1,A2,A3,B1,B2,B3,C1,C2,C3"', allow_blank=True)
    ws.add_data_validation(dv_sales); dv_sales.add("D4:D33")

    # 条件格式
    type_col = "F"
    for code, color in [("A", "FFC000"), ("B", "FF8C00"),
                         ("C", "FF6B6B"), ("D", "9966CC"),
                         ("E", "FFB6C1"), ("F", "808080"), ("G", "5B9BD5")]:
        ws.conditional_formatting.add(f"{type_col}4:{type_col}33",
            CellIsRule(operator="containsText", formula=[f'"{code}"'],
                       fill=PatternFill("solid", fgColor=color)))

    # 示例
    sample = ["A", "邯郸 XX 钢贸", "B1", "2026-06-20", "B价格",
              3, 2, "竞品价格便宜 40%", "没讲清楚价值故事", "2026-06-22",
              "你们贵但我没看到值多少", "降低期望换签"]
    for j, v in enumerate(sample, 2):
        ws.cell(row=4, column=j, value=v)

    ws.freeze_panes = "D4"


def build_type_dist(ws):
    ws.title = "7大类型分布"
    widths(ws, [4, 14, 12, 12, 30])

    ws.merge_cells("B2:E2")
    ws["B2"] = "7 大失败类型自动统计"
    ws["B2"].font = TITLE_FONT
    ws["B2"].alignment = LEFT

    headers = ["类型", "数量", "占比", "改进重点"]
    for j, h in enumerate(headers, 2):
        style_h(ws.cell(row=4, column=j, value=h))

    types = [
        ("A 时机", "A时机", "客户筛选 + 0 首付 + 6 月再启动"),
        ("B 价格", "B价格", "不正面降价 + 算 ROI + 三位一体捆绑"),
        ("C 价值", "C价值", "强制顾问式 + 5 客户故事 + 健康度自评"),
        ("D 关系", "D关系", "多层接触 + CRM 标签 + 季度回访"),
        ("E 信任", "E信任", "风险逆转 + 同行背书 + CS SOP"),
        ("F 竞品", "F竞品", "不正面打 + 差异化 + 未来再争夺"),
        ("G 内部", "G内部", "培养内部代言人 + 绕开反对者"),
    ]
    for i, (label, code, action) in enumerate(types, 5):
        bg = LIGHT_BG if i % 2 == 0 else WHITE
        style_t(ws.cell(row=i, column=2, value=label), fill=bg, bold=True)
        cell = ws.cell(row=i, column=3, value=f'=COUNTIF(\'失败客户登记表\'!F4:F33,"{code}")')
        style_t(cell, fill=SUCCESS, bold=True)
        cell = ws.cell(row=i, column=4,
                      value=f'=IFERROR(C{i}/SUM($C$5:$C$11),0)')
        style_t(cell, fill=SUCCESS, pct=True, bold=True)
        style_t(ws.cell(row=i, column=5, value=action), fill=bg, align=LEFT)

    # 合计
    i = 12
    style_t(ws.cell(row=i, column=2, value="合计"), fill=PRIMARY, bold=True)
    ws.cell(row=i, column=2).font = Font(name="Microsoft YaHei", size=11, bold=True, color=WHITE)
    cell = ws.cell(row=i, column=3, value="=SUM(C5:C11)")
    style_t(cell, fill=PRIMARY, bold=True)
    ws.cell(row=i, column=3).font = Font(name="Microsoft YaHei", size=11, bold=True, color=WHITE)


def build_sales_analysis(ws):
    ws.title = "销售个人失败分析"
    widths(ws, [4, 10, 12, 12, 12, 12, 12, 12, 12, 12, 30])

    ws.merge_cells("B2:K2")
    ws["B2"] = "9 销售个人失败客户分析"
    ws["B2"].font = TITLE_FONT
    ws["B2"].alignment = LEFT

    headers = ["销售", "总失败", "A时机", "B价格", "C价值", "D关系", "E信任", "F竞品", "G内部", "主要改进点"]
    for j, h in enumerate(headers, 2):
        style_h(ws.cell(row=4, column=j, value=h))

    sales = ["A1", "A2", "A3", "B1", "B2", "B3", "C1", "C2", "C3"]
    for i, name in enumerate(sales, 5):
        bg = LIGHT_BG if i % 2 == 0 else WHITE
        style_t(ws.cell(row=i, column=2, value=name), fill=bg, bold=True)
        # 总失败
        cell = ws.cell(row=i, column=3,
                      value=f'=COUNTIF(\'失败客户登记表\'!D4:D33,"{name}")')
        style_t(cell, fill=SUCCESS, bold=True)
        # 7 类失败数
        type_codes = ["A时机", "B价格", "C价值", "D关系", "E信任", "F竞品", "G内部"]
        for j, code in enumerate(type_codes, 4):
            cell = ws.cell(row=i, column=j,
                          value=(f'=COUNTIFS(\'失败客户登记表\'!D4:D33,"{name}",'
                                 f'\'失败客户登记表\'!F4:F33,"{code}")'))
            style_t(cell, fill=bg)
        # 主要改进点
        style_t(ws.cell(row=i, column=11), fill=INPUT_BG, align=LEFT)

    # 合计
    i = 14
    style_t(ws.cell(row=i, column=2, value="合计"), fill=PRIMARY, bold=True)
    ws.cell(row=i, column=2).font = Font(name="Microsoft YaHei", size=11, bold=True, color=WHITE)
    for col in range(3, 11):
        col_letter = get_column_letter(col)
        cell = ws.cell(row=i, column=col, value=f"=SUM({col_letter}5:{col_letter}13)")
        style_t(cell, fill=PRIMARY, bold=True)
        ws.cell(row=i, column=col).font = Font(name="Microsoft YaHei", size=11, bold=True, color=WHITE)


def build_product_analysis(ws):
    ws.title = "产品价格分析"
    widths(ws, [4, 16, 14, 14, 30])

    ws.merge_cells("B2:E2")
    ws["B2"] = "产品 / 价格失败分析"
    ws["B2"].font = TITLE_FONT
    ws["B2"].alignment = LEFT

    # 高频反对意见
    ws.merge_cells("B4:E4")
    style_t(ws.cell(row=4, column=2, value="▼ 高频反对意见 TOP 10"),
            fill=PRIMARY, align=LEFT, bold=True)
    ws.cell(row=4, column=2).font = Font(name="Microsoft YaHei", size=11, bold=True, color=WHITE)

    headers = ["排名", "反对意见", "出现次数", "改进对策"]
    for j, h in enumerate(headers, 2):
        style_h(ws.cell(row=5, column=j, value=h))

    objections = [
        ("竞品价格便宜", "", "出价值故事 + ROI 算法 + 三位一体捆绑"),
        ("我现在用别家", "", "强调互补 + 30 天免费试用"),
        ("没时间", "", "5 分钟演示 + 关键功能聚焦"),
        ("员工不会用", "", "10 分钟培训承诺 + 视频教学"),
        ("不相信 AI", "", "30 天免费 + 同行案例"),
        ("先问别人", "", "案例 PDF + 同行电话联系"),
        ("没听说过你们", "", "权威背书 + 行业排名 + 客户名单"),
        ("用别家 AI", "", "互补不替代"),
        ("看资料", "", "30 天免费试用"),
        ("找老板说", "", "先做『教练』，让他能讲清"),
    ]
    for i, (obj, count, action) in enumerate(objections, 6):
        style_t(ws.cell(row=i, column=2, value=i-5), fill=LIGHT_BG, bold=True)
        style_t(ws.cell(row=i, column=3, value=obj), fill=LIGHT_BG, align=LEFT)
        style_t(ws.cell(row=i, column=4, value=count), fill=INPUT_BG, bold=True)
        style_t(ws.cell(row=i, column=5, value=action), fill=LIGHT_BG, align=LEFT)


def build_improvement(ws):
    ws.title = "流程改进追踪"
    widths(ws, [4, 30, 14, 14, 14, 14, 22])

    ws.merge_cells("B2:G2")
    ws["B2"] = "流程改进动作追踪"
    ws["B2"].font = TITLE_FONT
    ws["B2"].alignment = LEFT

    headers = ["改进动作", "来源", "责任人", "开始日", "完成日", "状态"]
    for j, h in enumerate(headers, 2):
        style_h(ws.cell(row=4, column=j, value=h))

    # 30 行空白
    for i in range(5, 35):
        for col in range(2, 8):
            style_t(ws.cell(row=i, column=col), fill=INPUT_BG, align=LEFT)

    # 状态下拉
    dv_status = DataValidation(type="list",
        formula1='"待办,进行中,已完成,延期"', allow_blank=True)
    ws.add_data_validation(dv_status); dv_status.add("G5:G34")

    # 条件格式
    for status, color in [("已完成", "00B050"), ("进行中", "FFC000"),
                          ("延期", "FF6B6B")]:
        ws.conditional_formatting.add("G5:G34",
            CellIsRule(operator="equal", formula=[f'"{status}"'],
                       fill=PatternFill("solid", fgColor=color)))


def build_monthly(ws):
    ws.title = "月度复盘汇总"
    widths(ws, [4, 22, 14, 14, 14])

    ws.merge_cells("B2:E2")
    ws["B2"] = "月度失败客户复盘汇总"
    ws["B2"].font = TITLE_FONT
    ws["B2"].alignment = LEFT

    headers = ["指标", "M1", "M2", "M3"]
    for j, h in enumerate(headers, 2):
        style_h(ws.cell(row=4, column=j, value=h))

    metrics = [
        "新增失败客户数",
        "A 时机型",
        "B 价格型",
        "C 价值型",
        "D 关系型",
        "E 信任型",
        "F 竞品型",
        "G 内部型",
        "失败 / 接触客户比",
        "完成访谈数",
        "已启动改进数",
        "已完成改进数",
    ]
    for i, m in enumerate(metrics, 5):
        style_t(ws.cell(row=i, column=2, value=m), fill=LIGHT_BG, align=LEFT, bold=True)
        for col in range(3, 6):
            style_t(ws.cell(row=i, column=col), fill=INPUT_BG)

    # 月度三问
    ws.merge_cells("B19:E19")
    style_t(ws.cell(row=19, column=2, value="▼ 月度三问"),
            fill=PRIMARY, align=LEFT, bold=True)
    ws.cell(row=19, column=2).font = Font(name="Microsoft YaHei", size=11, bold=True, color=WHITE)

    questions = [
        "1. 本月失败的最大共性是什么？",
        "2. 哪个销售改进最快 / 最慢？",
        "3. 下月最重要的 1 个改进是？",
    ]
    for i, q in enumerate(questions, 20):
        style_t(ws.cell(row=i, column=2, value=q), fill=LIGHT_BG, align=LEFT, bold=True)
        ws.merge_cells(start_row=i, start_column=3, end_row=i, end_column=5)
        style_t(ws.cell(row=i, column=3), fill=INPUT_BG, align=LEFT)
        ws.row_dimensions[i].height = 40


def main():
    wb = Workbook()
    build_readme(wb.active)
    build_registry(wb.create_sheet())
    build_type_dist(wb.create_sheet())
    build_sales_analysis(wb.create_sheet())
    build_product_analysis(wb.create_sheet())
    build_improvement(wb.create_sheet())
    build_monthly(wb.create_sheet())

    out = "tools/失败客户分析模板.xlsx"
    wb.save(out)
    print(f"已生成：{out}")
    print(f"工作表：{wb.sheetnames}")


if __name__ == "__main__":
    main()
