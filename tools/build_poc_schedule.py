"""
灯塔厂 30 天 POC 部署日程表
============================

输出：tools/灯塔厂30天POC_部署清单.xlsx

工作表：
    1. 使用说明
    2. 30 天日程总览        —— 4 周关键节点
    3. 每日任务清单          —— D1-D30 每天动作 + 责任人 + 状态
    4. KPI 进度跟踪          —— 4 个核心 KPI 的 D7/D14/D21/D30 检查
    5. 风险预警表
    6. 30 天预算明细
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule


THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PRIMARY = "1F4E79"
ACCENT = "2E75B6"
LIGHT_BG = "DEEBF7"
INPUT_BG = "FFF2CC"
SUCCESS = "E2EFDA"
WARN = "FCE4D6"
CRITICAL = "FFB6C1"
WHITE = "FFFFFF"

H_FILL = PatternFill("solid", fgColor=PRIMARY)
H_FONT = Font(name="Microsoft YaHei", size=11, bold=True, color=WHITE)
TITLE_FONT = Font(name="Microsoft YaHei", size=16, bold=True, color=PRIMARY)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def style_h(c):
    c.fill = H_FILL; c.font = H_FONT; c.alignment = CENTER; c.border = BORDER


def style_t(c, fill=None, align=None, money=False, bold=False):
    c.font = Font(name="Microsoft YaHei", size=10, bold=bold)
    c.alignment = align or CENTER; c.border = BORDER
    if fill: c.fill = PatternFill("solid", fgColor=fill)
    if money: c.number_format = '"¥"#,##0'


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_readme(ws):
    ws.title = "使用说明"
    widths(ws, [4, 26, 70])
    ws.merge_cells("B2:C2")
    ws["B2"] = "灯塔厂 30 天 POC 部署清单（v1.0）"
    ws["B2"].font = TITLE_FONT
    ws["B2"].alignment = LEFT

    rows = [
        ("适用对象", "CTO（亲自跟）+ 实施 + 工程师 + 销售总监"),
        ("核心信条", "第 1 家灯塔厂的成败决定后续 5 家"),
        ("4 周节奏",
         "第 1 周（D1-D7）：部署 + 启动\n"
         "第 2 周（D8-D14）：数据见效 + 拍视频\n"
         "第 3 周（D15-D21）：KPI 冲刺验收\n"
         "第 4 周（D22-D30）：终验 + 营销转化"),
        ("4 个 KPI",
         "1. 卷号追溯响应 ≤ 5 秒\n"
         "2. 系统数据上传成功率 ≥ 99%\n"
         "3. 用户日活率 ≥ 80%\n"
         "4. 客户老板 NPS ≥ 8/10"),
        ("CTO 投入",
         "D1-D3 80% / D4-D7 30% / D8-D14 10% / D15-D21 30% / D22-D30 15%"),
        ("使用步骤",
         "① 打开『每日任务清单』，每天填『完成情况』列\n"
         "② D14 切到『KPI 进度跟踪』做中期评估\n"
         "③ D21 看 KPI 是否全达成 → 签 POC 验收单\n"
         "④ D30 终验 + 灯塔厂授牌"),
    ]
    for i, (k, v) in enumerate(rows, start=4):
        c1 = ws.cell(row=i, column=2, value=k); style_t(c1, fill=PRIMARY, bold=True)
        c1.font = Font(name="Microsoft YaHei", size=11, bold=True, color=WHITE)
        c2 = ws.cell(row=i, column=3, value=v); style_t(c2, align=LEFT)
        ws.row_dimensions[i].height = max(50, len(v) * 0.7)


def build_overview(ws):
    ws.title = "30天日程总览"
    widths(ws, [4, 12, 20, 30, 18, 14])
    ws.merge_cells("B2:F2")
    ws["B2"] = "30 天日程总览（4 周关键节点）"
    ws["B2"].font = TITLE_FONT; ws["B2"].alignment = LEFT

    headers = ["周", "天数", "主题", "关键里程碑", "CTO 投入", "状态"]
    for j, h in enumerate(headers, 2):
        style_h(ws.cell(row=4, column=j, value=h))

    weeks = [
        ("第 1 周", "D1-D7", "部署 + 启动",
         "硬件入场 / 软件部署 / 操作工培训 / 第 1 周数据小结",
         "80% → 30%", ""),
        ("第 2 周", "D8-D14", "数据见效 + 拍视频",
         "数据稳定 / 老板拍 3 分钟视频 / 第 1 篇公众号上线",
         "10%", ""),
        ("第 3 周", "D15-D21", "KPI 冲刺验收",
         "性能优化 / D20 验收预演 / D21 正式 KPI 验收",
         "30%", ""),
        ("第 4 周", "D22-D30", "终验 + 营销转化",
         "案例打磨 / 经营版升级谈判 / D30 终验 + 授牌",
         "15%", ""),
    ]
    colors = [LIGHT_BG, SUCCESS, WARN, "FFE699"]
    for i, (w, days, theme, ms, cto, status) in enumerate(weeks, 5):
        bg = colors[i - 5]
        style_t(ws.cell(row=i, column=2, value=w), fill=bg, bold=True)
        style_t(ws.cell(row=i, column=3, value=days), fill=bg, bold=True)
        style_t(ws.cell(row=i, column=4, value=theme), fill=bg, align=LEFT)
        style_t(ws.cell(row=i, column=5, value=ms), fill=bg, align=LEFT)
        style_t(ws.cell(row=i, column=6, value=cto), fill=bg)
        style_t(ws.cell(row=i, column=7, value=status), fill=INPUT_BG)
        ws.row_dimensions[i].height = 50


def build_daily(ws):
    ws.title = "每日任务清单"
    widths(ws, [4, 8, 14, 36, 16, 16, 18])
    ws.merge_cells("B2:G2")
    ws["B2"] = "30 天每日任务清单（黄色 = 状态 / 备注 输入）"
    ws["B2"].font = TITLE_FONT; ws["B2"].alignment = LEFT

    headers = ["天", "周次", "动作", "责任人", "完成状态", "备注"]
    for j, h in enumerate(headers, 2):
        style_h(ws.cell(row=4, column=j, value=h))

    # 30 天任务列表
    tasks = [
        # (天, 周, 动作, 责任人)
        ("D1", "W1", "客户启动会（老板+副总+班长齐聚）", "CTO+客户老板"),
        ("D1", "W1", "现场踩点（产线/仓库/网络）", "CTO+实施"),
        ("D1", "W1", "硬件安装：标签机+扫码枪+平板", "实施+工程师"),
        ("D2", "W1", "服务器/SaaS 账号开通+基础数据初始化", "工程师"),
        ("D2", "W1", "客户 IT/财务导入 5 张 Excel", "客户 IT+实施"),
        ("D2", "W1", "后台跑通第 1 张报表", "工程师"),
        ("D3", "W1", "白班操作工 1 小时培训（扫码+平板）", "实施"),
        ("D3", "W1", "晚班操作工 1 小时培训", "实施"),
        ("D3", "W1", "班长 30 分钟培训", "实施"),
        ("D4-D6", "W1", "实施工程师驻场全程跟扫码", "实施+工程师"),
        ("D4-D6", "W1", "每天 22:00 给客户老板发数据简报微信", "实施"),
        ("D7", "W1", "★ 第 1 周数据小结 + 客户老板第 1 份正式报告", "CTO+实施"),
        # W2
        ("D8-D11", "W2", "现场频次降到每 2 天 1 次（不再驻场）", "实施"),
        ("D8-D11", "W2", "远程监控数据上传率 + 异常处理", "工程师"),
        ("D8-D11", "W2", "持续每天发数据简报", "实施"),
        ("D12-D13", "W2", "与客户老板沟通拍视频意向 + 5 句话脚本", "老板+市场"),
        ("D12-D13", "W2", "拍摄团队对接：拍什么 / 角度 / 道具", "市场"),
        ("D14", "W2", "★ 拍客户老板 3 分钟视频 + 卷号 5 秒查到演示画面", "市场"),
        ("D14", "W2", "★ 第 1 篇公众号 + 1 条短视频上线", "市场"),
        # W3
        ("D15-D19", "W3", "系统性能优化（缓存/索引/标签机改造）", "工程师"),
        ("D15-D19", "W3", "推动客户用第 2 个查询场景", "实施+客户副总"),
        ("D15-D19", "W3", "客户老板亲自试一次手机查询（关键！）", "老板+客户老板"),
        ("D20", "W3", "★ 验收预演：跑 3 个测试场景", "CTO+实施+工程师"),
        ("D20", "W3", "输出《验收预演报告》", "实施"),
        ("D21", "W3", "★ 正式 KPI 验收（客户老板+CTO+销售总监现场）", "CTO+老板+客户老板"),
        ("D21", "W3", "签署《POC 验收单》", "全员"),
        ("D21", "W3", "拍验收照片 + 视频", "市场"),
        # W4
        ("D22-D26", "W4", "21 天数据做成《灯塔厂案例报告》", "实施+市场"),
        ("D22-D26", "W4", "打磨第 2 篇公众号「30 天数据故事」", "市场"),
        ("D22-D26", "W4", "准备灯塔厂老板老板局演讲 30 分钟稿", "老板+市场"),
        ("D22-D26", "W4", "协会/行业媒体接洽（CCTV-2/找钢网/Mysteel）", "市场+老板"),
        ("D27-D29", "W4", "★ 客户老板录第 2 段视频「30 天我们做对了什么」", "市场"),
        ("D27-D29", "W4", "与客户讨论经营版升级 + 销售总监正式报价", "销售总监+老板"),
        ("D30", "W4", "★ 终验大会：客户+我方 CEO+协会代表（如可邀）", "全员"),
        ("D30", "W4", "★ 灯塔厂正式授牌（牌匾装裱）", "老板"),
        ("D30", "W4", "拍授牌合影 + 团队大合影", "市场"),
        ("D30", "W4", "公司战报朋友圈 + 全员庆功", "全员"),
    ]

    # 状态下拉
    dv_status = DataValidation(type="list",
        formula1='"未开始,进行中,已完成,延期,放弃"', allow_blank=True)

    week_colors = {"W1": LIGHT_BG, "W2": SUCCESS, "W3": WARN, "W4": "FFE699"}
    for i, (day, week, action, owner) in enumerate(tasks, 5):
        bg = week_colors.get(week, WHITE)
        style_t(ws.cell(row=i, column=2, value=day), fill=bg, bold=True)
        style_t(ws.cell(row=i, column=3, value=week), fill=bg)
        style_t(ws.cell(row=i, column=4, value=action), fill=bg, align=LEFT)
        style_t(ws.cell(row=i, column=5, value=owner), fill=bg)
        style_t(ws.cell(row=i, column=6, value=""), fill=INPUT_BG)
        style_t(ws.cell(row=i, column=7, value=""), fill=INPUT_BG)
        ws.row_dimensions[i].height = 22

    ws.add_data_validation(dv_status)
    dv_status.add(f"F5:F{4+len(tasks)}")

    # 条件格式
    for col in ["F"]:
        ws.conditional_formatting.add(f"{col}5:{col}{4+len(tasks)}",
            CellIsRule(operator="equal", formula=['"已完成"'],
                       fill=PatternFill("solid", fgColor="00B050")))
        ws.conditional_formatting.add(f"{col}5:{col}{4+len(tasks)}",
            CellIsRule(operator="equal", formula=['"延期"'],
                       fill=PatternFill("solid", fgColor="FFC000")))
        ws.conditional_formatting.add(f"{col}5:{col}{4+len(tasks)}",
            CellIsRule(operator="equal", formula=['"放弃"'],
                       fill=PatternFill("solid", fgColor="FF6B6B")))

    ws.freeze_panes = "B5"


def build_kpi(ws):
    ws.title = "KPI进度跟踪"
    widths(ws, [4, 26, 14, 14, 14, 14, 14, 22])
    ws.merge_cells("B2:H2")
    ws["B2"] = "4 个核心 KPI 在 D7/D14/D21/D30 的进度跟踪"
    ws["B2"].font = TITLE_FONT; ws["B2"].alignment = LEFT

    headers = ["KPI", "目标", "D7 实际", "D14 实际", "D21 实际", "D30 实际", "D21 是否达成"]
    for j, h in enumerate(headers, 2):
        style_h(ws.cell(row=4, column=j, value=h))

    kpis = [
        ("卷号追溯响应（秒）", "≤ 5", "≤ 30 秒可", "≤ 10 秒", "≤ 5 秒（必达）", "保持"),
        ("系统数据上传成功率", "≥ 99%", "≥ 95%", "≥ 98%", "≥ 99%（必达）", "≥ 99%"),
        ("用户日活率", "≥ 80%", "≥ 60%", "≥ 75%", "≥ 80%（必达）", "≥ 85%"),
        ("客户老板 NPS（10 分制）", "≥ 8", "≥ 7", "≥ 7", "≥ 8（必达）", "≥ 9"),
    ]
    for i, (kpi, target, d7, d14, d21, d30) in enumerate(kpis, 5):
        style_t(ws.cell(row=i, column=2, value=kpi), fill=LIGHT_BG, bold=True, align=LEFT)
        style_t(ws.cell(row=i, column=3, value=target), fill=SUCCESS, bold=True)
        # D7 D14 D21 D30 输入
        for col in range(4, 8):
            style_t(ws.cell(row=i, column=col, value=""), fill=INPUT_BG)
        # D21 达成
        style_t(ws.cell(row=i, column=8, value=""), fill=INPUT_BG, bold=True)

    # D14 中期评估
    ws.merge_cells("B11:H11")
    ws["B11"] = "★ D14 中期评估（任 1 项红灯 → D15-D21 改驻场模式）"
    style_t(ws["B11"], fill=PRIMARY, align=LEFT, bold=True)
    ws["B11"].font = Font(name="Microsoft YaHei", size=11, bold=True, color=WHITE)

    eval_headers = ["检查项", "标准", "实际", "状态", "红灯怎么办", ""]
    for j, h in enumerate(eval_headers[:5], 2):
        style_h(ws.cell(row=12, column=j, value=h))
    evals = [
        ("扫码完成率",       "≥ 80%",   "", "", "< 60% 立刻加现场资源"),
        ("客户老板满意度",   "≥ 7/10",  "", "", "< 6 立刻 CTO 上门聊"),
        ("系统稳定性",       "数据丢失 < 1%", "", "", "> 5% 立刻硬件升级"),
        ("客户副总反馈",     "主动 1-2 次", "", "", "0 次拉到 IM 群里"),
    ]
    for i, (item, target, real, status, redaction) in enumerate(evals, 13):
        style_t(ws.cell(row=i, column=2, value=item), fill=LIGHT_BG, align=LEFT, bold=True)
        style_t(ws.cell(row=i, column=3, value=target), fill=SUCCESS, bold=True)
        style_t(ws.cell(row=i, column=4, value=real), fill=INPUT_BG)
        style_t(ws.cell(row=i, column=5, value=status), fill=INPUT_BG)
        ws.merge_cells(start_row=i, start_column=6, end_row=i, end_column=7)
        style_t(ws.cell(row=i, column=6, value=redaction), align=LEFT)


def build_risk(ws):
    ws.title = "风险预警"
    widths(ws, [4, 24, 22, 30])
    ws.merge_cells("B2:D2")
    ws["B2"] = "30 天部署 8 大风险预警"
    ws["B2"].font = TITLE_FONT; ws["B2"].alignment = LEFT

    headers = ["风险", "触发信号", "预案"]
    for j, h in enumerate(headers, 2):
        style_h(ws.cell(row=4, column=j, value=h))

    risks = [
        ("工人抵触扫码", "D3-D5 扫码完成率 < 50%",
         "现场加奖励（每条扫卷 ¥1）+ 班长沟通"),
        ("老员工『我习惯老办法』", "D7 仍有不配合",
         "客户老板亲自训话 + 不扫不发工资"),
        ("网络不稳定", "数据丢失 / 上传失败",
         "加 4G 路由器 + 本地缓存机制"),
        ("客户副总配合度低", "7 天没主动联系",
         "我方销售总监找客户老板复述重要性"),
        ("客户老板观望", "不主动看报表",
         "主动每天发数据 + 邀请到现场看"),
        ("数据偏差大（成本类）", "4 条线对比数字失真",
         "检查分摊规则 + 让财务复核"),
        ("标签机/扫码枪故障", "第 2 周后频发",
         "备用机随时调换；故障率 > 5% 整批换"),
        ("KPI 不达标", "D21 验收失败",
         "启动『3 天加速包』（CTO + 工程师驻场）"),
    ]
    for i, (risk, signal, plan) in enumerate(risks, 5):
        bg = LIGHT_BG if i % 2 == 0 else WHITE
        style_t(ws.cell(row=i, column=2, value=risk), fill=bg, align=LEFT, bold=True)
        style_t(ws.cell(row=i, column=3, value=signal), fill=bg, align=LEFT)
        style_t(ws.cell(row=i, column=4, value=plan), fill=bg, align=LEFT)
        ws.row_dimensions[i].height = 36


def build_budget(ws):
    ws.title = "30天预算"
    widths(ws, [4, 30, 14, 40])
    ws.merge_cells("B2:D2")
    ws["B2"] = "30 天部署预算明细"
    ws["B2"].font = TITLE_FONT; ws["B2"].alignment = LEFT

    headers = ["项", "金额（元）", "备注"]
    for j, h in enumerate(headers, 2):
        style_h(ws.cell(row=4, column=j, value=h))

    items = [
        ("硬件（标签机+扫码枪×2+平板×1）", 5000, "灯塔厂送 / 试用结束可回收"),
        ("CTO + 实施 + 工程师差旅", 15000, "30 天累计"),
        ("培训物料 + 礼品", 2000, "培训简卡 + 钢制纪念品"),
        ("视频拍摄 + 剪辑", 8000, "2 段视频 + 1 篇公众号配图"),
        ("授牌 + 牌匾 + 合影", 3000, "终验仪式"),
        ("应急 / 加班费", 5000, ""),
    ]
    total = 0
    for i, (name, amount, note) in enumerate(items, 5):
        style_t(ws.cell(row=i, column=2, value=name), fill=LIGHT_BG, align=LEFT, bold=True)
        style_t(ws.cell(row=i, column=3, value=amount), fill=SUCCESS, money=True, bold=True)
        style_t(ws.cell(row=i, column=4, value=note), align=LEFT)
        total += amount

    # 合计
    i = 5 + len(items)
    style_t(ws.cell(row=i, column=2, value="灯塔厂部署小计"), fill=PRIMARY)
    ws.cell(row=i, column=2).font = Font(name="Microsoft YaHei", size=11,
                                          bold=True, color=WHITE)
    style_t(ws.cell(row=i, column=3, value=total), fill=PRIMARY, money=True, bold=True)
    ws.cell(row=i, column=3).font = Font(name="Microsoft YaHei", size=11,
                                          bold=True, color=WHITE)

    # 加上灯塔厂实施费打折
    i += 1
    style_t(ws.cell(row=i, column=2, value="+ 灯塔厂实施费打折部分"),
            fill=LIGHT_BG, align=LEFT, bold=True)
    style_t(ws.cell(row=i, column=3, value=30000), fill=WARN, money=True, bold=True)
    style_t(ws.cell(row=i, column=4, value="实施费 5 折损失约 ¥3 万"), align=LEFT)

    i += 1
    style_t(ws.cell(row=i, column=2, value="单家灯塔厂总投入"), fill=PRIMARY)
    ws.cell(row=i, column=2).font = Font(name="Microsoft YaHei", size=11,
                                          bold=True, color=WHITE)
    style_t(ws.cell(row=i, column=3, value=total + 30000),
            fill=PRIMARY, money=True, bold=True)
    ws.cell(row=i, column=3).font = Font(name="Microsoft YaHei", size=11,
                                          bold=True, color=WHITE)

    # ROI 说明
    i += 2
    ws.merge_cells(start_row=i, start_column=2, end_row=i+2, end_column=4)
    ws.cell(row=i, column=2).value = (
        "★ ROI 说明：\n"
        "1 家灯塔厂带来直接收入 ¥10-25 万 + 间接收入 ¥50-150 万 + 行业地位\n"
        "灯塔厂 ROI = 7-25 倍，比任何营销渠道都高"
    )
    ws.cell(row=i, column=2).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws.cell(row=i, column=2).font = Font(name="Microsoft YaHei", size=10, color="595959")
    ws.cell(row=i, column=2).fill = PatternFill("solid", fgColor=SUCCESS)
    ws.cell(row=i, column=2).border = BORDER


def main():
    wb = Workbook()
    build_readme(wb.active)
    build_overview(wb.create_sheet())
    build_daily(wb.create_sheet())
    build_kpi(wb.create_sheet())
    build_risk(wb.create_sheet())
    build_budget(wb.create_sheet())

    out = "tools/灯塔厂30天POC_部署清单.xlsx"
    wb.save(out)
    print(f"已生成：{out}")
    print(f"工作表：{wb.sheetnames}")


if __name__ == "__main__":
    main()
